# -*- coding: utf-8 -*-
"""通讯录表格解析：CSV / Excel(xlsx) → [{name, email, aliases}]。

灵活识别表头：First/Last Name + Email，或单个 姓名/Name + 邮箱/Email。
"""
from __future__ import annotations

import csv
import io


def parse_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def parse_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    return [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]


def rows_to_entries(rows: list[list[str]]) -> list[dict]:
    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if not rows:
        return []
    header = [(h or "").strip().lower() for h in rows[0]]

    def find(*keys):
        for i, h in enumerate(header):
            if any(k in h for k in keys):
                return i
        return -1

    i_email = find("email", "邮箱", "e-mail", "mail")
    i_first = find("first name", "名字", "given name")
    i_last = find("last name", "姓氏", "family name")
    i_full = find("full name", "姓名", "name") if i_first < 0 else -1

    if i_email < 0:  # 邮箱列找不到：取"值里含@最多"的列
        best, best_n = -1, 0
        for ci in range(len(header)):
            n = sum(1 for r in rows[1:] if "@" in (r[ci] if ci < len(r) else ""))
            if n > best_n:
                best, best_n = ci, n
        i_email = best

    entries = []
    for r in rows[1:]:
        def cell(i):
            return (r[i].strip() if 0 <= i < len(r) and r[i] else "")
        email = cell(i_email)
        if not email or "@" not in email:
            continue
        if i_first >= 0 or i_last >= 0:
            first, last = cell(i_first), cell(i_last)
            name = f"{first} {last}".strip()
            aliases = [x for x in (first, last) if x]
        else:
            name = cell(i_full) or email.split("@")[0]
            aliases = []
        entries.append({"name": name or email.split("@")[0], "email": email, "aliases": aliases})
    return entries


def parse_table(content: bytes, filename: str = "") -> list[dict]:
    """按内容/文件名判断 CSV 还是 xlsx，解析成通讯录条目。"""
    is_xlsx = content[:4] == b"PK\x03\x04" or filename.lower().endswith((".xlsx", ".xls"))
    rows = parse_xlsx(content) if is_xlsx else parse_csv(content)
    return rows_to_entries(rows)
